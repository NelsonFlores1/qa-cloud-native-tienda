from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


BASE = Path(r"C:\Users\Nelson\Documents\GitHub\qa-cloud-native-tienda\docs\documento-soporte-reto-modulo5_v2_base.docx")
OUT = Path(r"C:\Users\Nelson\Documents\GitHub\qa-cloud-native-tienda\docs\documento-soporte-reto-modulo5_v2_fase6.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def apply_cell_text(cell, text, bold=False, color=None, size=9):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        set_cell_shading(cell, "1F4E78")
        apply_cell_text(cell, header, bold=True, color=(255, 255, 255), size=8.5)
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].width = widths[index]
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "EAF2F8")
            apply_cell_text(cells[index], value, size=8.5)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
    doc.add_paragraph()
    return table


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    return p


def main():
    doc = Document(BASE)
    headings = [p for p in doc.paragraphs if p.text.strip() == "Fase 6 Metricas y estrategia"]
    if not headings:
        doc.add_heading("Fase 6 Metricas y estrategia", level=1)

    add_body(
        doc,
        "La fase 6 consolida la evidencia obtenida en el reto y define una linea base de mejora continua. "
        "Las metricas se interpretan como indicadores academicos del flujo CI/CD: el proyecto no realiza despliegues "
        "a produccion y el stage Deploy permanece simulado."
    )

    doc.add_heading("Indicadores DORA", level=2)
    add_body(
        doc,
        "Los indicadores DORA permiten observar la velocidad y estabilidad del flujo. Para no presentar datos "
        "de produccion inexistentes como si fueran reales, se separan la evidencia de clase, la forma de medicion "
        "y las metas propuestas."
    )
    add_table(
        doc,
        ["Indicador", "Linea base del reto", "Como se medira", "Meta inicial"],
        [
            (
                "Frecuencia de entrega",
                "Una imagen tienda-qa publicada en ACR con el tag 7. El Deploy es simulado.",
                "Conteo de publicaciones exitosas en ACR por semana.",
                "Al menos una publicacion por semana cuando haya cambios aprobados.",
            ),
            (
                "Lead time for changes",
                "La ejecucion completa con gates verdes tomo aproximadamente 7 min 37 s.",
                "Tiempo entre el commit y la publicacion verde de la imagen.",
                "Menor o igual a 10 minutos para cambios de bajo riesgo.",
            ),
            (
                "Tasa de fallos de cambio",
                "Se detectaron dos fallos de configuracion en fase 5: dependencia Express ausente y ACR Tasks restringido.",
                "Cambios que requieren correccion despues de CI dividido entre cambios publicados.",
                "Menor o igual a 15 por ciento; excluir ejercicios de aprendizaje del indicador de produccion.",
            ),
            (
                "MTTR",
                "No hubo incidentes de produccion. Los fallos de CI se corrigieron antes de publicar la imagen.",
                "Tiempo desde el primer pipeline rojo hasta la primera corrida verde posterior.",
                "Registrar cada recuperacion y reducir el tiempo en iteraciones posteriores.",
            ),
        ],
        [Pt(78), Pt(155), Pt(145), Pt(128)],
    )

    doc.add_heading("Piramide de pruebas actual", level=2)
    add_body(
        doc,
        "La suite valida el flujo de mayor valor mediante pruebas de humo, carga, E2E y contrato. La principal "
        "oportunidad de mejora es aumentar la base de pruebas unitarias para reducir la dependencia de pruebas E2E "
        "mas lentas y sensibles a la interfaz."
    )
    add_table(
        doc,
        ["Nivel", "Evidencia implementada", "Cobertura y decision"],
        [
            (
                "Unitarias",
                "No hay una suite unitaria dedicada.",
                "Agregar pruebas al generador de datos y a las reglas de rutas API antes de ampliar la interfaz.",
            ),
            (
                "API y contrato",
                "Smoke API y un contrato Pact para GET /api/productos.",
                "Valida disponibilidad, respuesta funcional y compatibilidad consumidor-proveedor.",
            ),
            (
                "E2E",
                "Siete escenarios pytest-bdd con Playwright y Screenplay.",
                "Cubre catalogo, carrito, pedido y tres casos de borde; se mantiene para rutas criticas.",
            ),
            (
                "No funcional",
                "Gate de carga con SLA de p95 y tasa de error.",
                "Evita promover cambios que degraden el rendimiento acordado.",
            ),
        ],
        [Pt(84), Pt(195), Pt(227)],
    )

    doc.add_heading("Riesgos y mitigaciones", level=2)
    add_table(
        doc,
        ["Riesgo", "Impacto", "Mitigacion aplicada o propuesta"],
        [
            (
                "Suscripcion educativa bloquea ACR Tasks.",
                "No se podia usar az acr build para publicar la imagen.",
                "El pipeline construye en el agente y publica con docker push usando la conexion sc-tienda-acr y el rol AcrPush.",
            ),
            (
                "Ejecucion cloud de navegadores sin cuenta Sauce Labs.",
                "No existe evidencia de una sesion remota real.",
                "Se conserva la demo guiada con capabilities, variables de entorno y stage opcional; para ejecucion real se requieren credenciales y una URL accesible.",
            ),
            (
                "Cobertura concentrada en E2E.",
                "Los diagnosticos pueden ser mas lentos y fragiles ante cambios visuales.",
                "Priorizar pruebas unitarias y de API; reservar E2E para escenarios de negocio de alto valor.",
            ),
            (
                "Deploy simulado.",
                "No se mide disponibilidad ni recuperacion de una aplicacion productiva.",
                "Como siguiente paso, desplegar una version etiquetada en un ambiente de prueba y medir salud y rollback.",
            ),
        ],
        [Pt(125), Pt(155), Pt(226)],
    )

    doc.add_heading("Estrategia de mejora continua", level=2)
    add_body(
        doc,
        "Cada cambio debe mantener la generacion reproducible de datos, superar los cinco gates de calidad y publicar "
        "un artefacto identificable en ACR. El equipo debe revisar semanalmente la duracion del pipeline, las fallas "
        "que bloquearon cambios y la evolucion de los indicadores DORA. Cuando exista un ambiente real, el gate de "
        "despliegue debe incluir smoke posterior al deploy, observabilidad y una estrategia de rollback."
    )
    add_body(
        doc,
        "Resultado de la practica: los gates funcionales, de carga, E2E, contenedor y contrato quedaron validados; "
        "la imagen tienda-qa se publico en Azure Container Registry. La fase cloud de navegadores se mantiene como "
        "demo guiada hasta contar con credenciales de Sauce Labs."
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
