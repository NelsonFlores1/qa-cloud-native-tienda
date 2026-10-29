from pathlib import Path
import sys
from docx import Document


source = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Users\Nelson\Documents\GitHub\qa-cloud-native-tienda\docs\documento-soporte-reto-modulo5_v2_base.docx")
doc = Document(source)

for index, paragraph in enumerate(doc.paragraphs, start=1):
    text = paragraph.text.strip().replace("\t", " ")
    if text:
        print(f"P {index} [{paragraph.style.name}] {text}")

for table_index, table in enumerate(doc.tables, start=1):
    print(f"TABLE {table_index}: {len(table.rows)}x{len(table.columns)}")
    for row in table.rows:
        print(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
