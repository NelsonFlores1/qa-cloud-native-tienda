from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\Nelson\Downloads\documento-soporte-reto-modulo5_v2.docx")
OUTPUT = Path(r"C:\Users\Nelson\Documents\GitHub\qa-cloud-native-tienda\docs\documento-soporte-reto-modulo5_v3.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def apply_run_font(run, size=10.0, bold=False, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def append_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    repeat_header(table.rows[0])

    for col, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[col]
        set_cell_width(cell, width)
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        apply_run_font(run, size=9.5, bold=True, color="FFFFFF")

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col, (value, width) in enumerate(zip(row, widths)):
            cell = cells[col]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, "EAF2F8")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if col == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            apply_run_font(run, size=9.3)

    doc.add_paragraph()
    return table


def set_paragraph(paragraph, text, style_name):
    paragraph.clear()
    paragraph.style = style_name
    run = paragraph.add_run(text)
    apply_run_font(run, size=11 if style_name == "Normal" else 12, bold=style_name.startswith("Heading"))


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 2")
    run = paragraph.add_run(text)
    apply_run_font(run, size=12, bold=True)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    apply_run_font(run, size=11)
    return paragraph


def main():
    doc = Document(SOURCE)

    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            if "Documento de soporte" in paragraph.text:
                paragraph.clear()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(
                    "QA Cloud Native Tienda  |  Documento de soporte  |  29 de octubre de 2026"
                )
                apply_run_font(run, size=9, color="808080")

    # The original document describes the first version of the pipeline. Keep the
    # section but make it consistent with the stages now verified in Azure DevOps.
    set_paragraph(
        doc.paragraphs[11],
        "El archivo pipeline/azure-pipelines.yml define un flujo con Build, SmokeTests, LoadTest, E2E, CloudE2E, Container, ContractTests, AcrPublication y Deploy simulado. Los gates verifican disponibilidad, SLA, escenarios BDD, salud del contenedor y contrato Pact. La publicacion en ACR se habilita manualmente y solo se realiza despues de los gates anteriores.",
        "Normal",
    )

    # The source document reserves two empty paragraphs immediately after phase 6.
    set_paragraph(doc.paragraphs[37], "Linea base academica DORA", "Heading 2")
    set_paragraph(
        doc.paragraphs[38],
        "Las siguientes metricas se presentan como linea base de la practica. La aplicacion no cuenta con trafico ni incidentes de produccion, por lo que se diferencian los hechos observados en Azure DevOps de los indicadores que deben medirse cuando exista un entorno productivo.",
        "Normal",
    )

    append_table(
        doc,
        ["Metrica", "Linea base observada", "Uso y siguiente medicion"],
        [
            (
                "Frecuencia de despliegue",
                "Una imagen publicada y verificada en ACR: repositorio tienda-qa, tag 7.",
                "Registrar una publicacion por cada cambio que supere todos los gates.",
            ),
            (
                "Lead time de cambio",
                "No hay serie historica de produccion. El pipeline deja el identificador de build y sus tiempos por etapa.",
                "Medir desde el commit hasta la publicacion del tag en ACR.",
            ),
            (
                "Tasa de fallos de cambio",
                "No existe despliegue productivo ni incidente real; los fallos de CI se resolvieron antes de la corrida final verde.",
                "Separar fallos de CI de fallos posteriores a produccion.",
            ),
            (
                "MTTR",
                "No aplica como metrica productiva en esta practica, porque no se registraron incidentes de produccion.",
                "Cuando exista operacion real, medir desde la deteccion hasta la restauracion del servicio.",
            ),
        ],
        [1.15, 2.85, 2.5],
    )

    add_heading(doc, "Piramide de pruebas y cobertura actual")
    add_body(
        doc,
        "La automatizacion actual concentra la mayor cantidad de escenarios en E2E. Es adecuada para demostrar el flujo completo del reto, pero la evolucion recomendada es aumentar pruebas unitarias para que los fallos se detecten antes y con menor tiempo de ejecucion.",
    )
    append_table(
        doc,
        ["Nivel", "Evidencia actual", "Decision de estrategia"],
        [
            (
                "Unitarias",
                "No se identifico una suite unitaria aislada para el generador o los modulos de la aplicacion.",
                "Agregar pruebas para generar-datos.js, validaciones de catalogo y calculos de carrito.",
            ),
            (
                "Contrato e integracion",
                "Pact valida GET /api/productos entre TiendaWeb y TiendaApi con consumidor y provider.",
                "Mantenerlo como gate antes de publicar la imagen.",
            ),
            (
                "E2E",
                "Siete escenarios Gherkin en cuatro features: catalogo, carrito, pedidos y bordes.",
                "Conservar casos criticos y mover reglas deterministas hacia pruebas unitarias.",
            ),
        ],
        [1.2, 2.9, 2.4],
    )

    add_heading(doc, "Riesgos y mitigaciones")
    append_table(
        doc,
        ["Riesgo", "Impacto", "Mitigacion aplicada o propuesta"],
        [
            (
                "Dependencia de servicios cloud",
                "Una restriccion de ACR Tasks puede bloquear la publicacion desde el pipeline.",
                "El pipeline usa Docker push y el principal de servicio tiene el rol AcrPush sobre el ACR.",
            ),
            (
                "Cobertura unitaria limitada",
                "Los fallos simples llegan a gates de mayor costo, como E2E.",
                "Priorizar una suite unitaria para datos, validaciones y logica de carrito.",
            ),
            (
                "Metricas sin produccion",
                "DORA no puede interpretarse como rendimiento operativo real.",
                "Etiquetar esta medicion como academica y registrar eventos reales en el siguiente entorno.",
            ),
            (
                "Ejecucion remota opcional",
                "Sauce Labs requiere cuenta, secretos y una URL accesible para una sesion real.",
                "Se conserva la demo guiada y se activara el stage CloudE2E al disponer de esas credenciales.",
            ),
        ],
        [1.45, 2.35, 2.7],
    )

    add_heading(doc, "Conclusion de la fase 6")
    add_body(
        doc,
        "La practica deja una linea de entrega verificable: datos reproducibles, gates de calidad, pruebas E2E, verificacion de contrato, healthcheck de contenedor e imagen publicada en ACR. La siguiente mejora prioritaria es crear pruebas unitarias y capturar de forma automatica los tiempos entre commit, publicacion y recuperacion para convertir la linea base academica en metricas DORA operativas.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
